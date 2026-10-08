"""日文課堂筆記 Word 產生器（實驗版）

讀取 note.json，產生仿原版講義版面的 .docx。
聲調線用「每拍一格」的表格做：上方小列放讀音並承載紅線，
高音拍畫上緣紅線，下降拍再畫右緣紅線；全部是 Word 原生表格框線，可編輯。

日文行的寫法（字串）：
  片段之間用「／」分隔。
  {漢字|讀音}  有讀音的漢字；其他假名直接寫。
  +助詞       接續文字（例：{割引|わりびき}+は），讀音不同時寫成 +は=わ。
  #n          該片段的聲調號數（省略時用 accent_default，null 表示待確認、不畫線）。
  !、 !。     標點（不畫線）。
  @A：        標籤（不畫線）。

區塊：title、tag、cards、bullets、defs、grid、sentences、note、pending，
欄位說明見 00_先讀我_知識總覽.md 第 7 節。
"""
import hashlib, json, os, re, sys, urllib.request
from docx import Document
from docx.shared import Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---- 樣式（數值取自原版 PDF） ----
FONT_JP = "Yu Gothic UI"          # 原版：黑體 UI（日本語）
FONT_TITLE = "源石黑體 M"
FONT_ZH = "清松手寫體1"
FONT_HEAD = "Calibri"
SZ_JP, SZ_RUBY, SZ_ZH, SZ_HEAD = 14, 8, 12, 10
BLUE = RGBColor(0x00, 0x70, 0xC0)
GRAY = RGBColor(0x76, 0x71, 0x71)
ORANGE = RGBColor(0xC5, 0x5A, 0x11)
PITCH_RED = "FF7D80"
BORDER_GRAY = "E8E6E6"
HEAD_FILL = "DEEAF6"
TAG_FILL = "FFE699"
PRACTICE_FILL = "E7E6E6"  # 短句練習表頭（灰）
CONTENT_W = 487  # pt，A4 扣左右邊界 54pt
ILLUST_W = 84    # pt，詞卡插圖寬（詞卡寬約 162pt）
IMG_CACHE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "images")

SMALL = set("ゃゅょぁぃぅぇぉゎャュョァィゥェォヮ")


# ---------- 小工具 ----------
def set_font(run, name, size, color=None, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    if color is not None:
        run.font.color.rgb = color


def tight(par, align=None):
    pf = par.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0
    if align is not None:
        par.alignment = align


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_borders(cell, **sides):
    """sides: top/right/bottom/left = (color, size_eighths) 或 None（無框線）"""
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for side in ("top", "left", "bottom", "right"):
        if side not in sides:
            continue
        el = OxmlElement(f"w:{side}")
        v = sides[side]
        if v is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(v[1]))
            el.set(qn("w:space"), "0"); el.set(qn("w:color"), v[0])
        b.append(el)


def table_borders(table, color=None, sz=8):
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        if color is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0"); el.set(qn("w:color"), color)
        b.append(el)
    tblPr.append(b)


def fixed_layout(table, widths_pt, cell_margin_pt=None):
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    if cell_margin_pt is not None:
        mar = OxmlElement("w:tblCellMar")
        for side in ("left", "right", "top", "bottom"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:w"), str(int(cell_margin_pt.get(side, 0) * 20))); el.set(qn("w:type"), "dxa")
            mar.append(el)
        tblPr.append(mar)
    tw = tblPr.find(qn("w:tblW"))  # 明確指定總寬，避免 LibreOffice 依內容縮放
    if tw is None:
        tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(int(sum(widths_pt) * 20)))
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_pt):
        gc.set(qn("w:w"), str(int(w * 20)))
    for row in table.rows:
        for c, w in zip(row.cells, widths_pt):
            c.width = Pt(w)


def first_par(cell):
    p = cell.paragraphs[0]
    tight(p)
    return p


def shrink_trailing(cell):
    """cell.add_table 會在表格後補一個空段落，把它縮到最小。"""
    p = cell.paragraphs[-1]
    tight(p)
    r = p.add_run(""); set_font(r, FONT_JP, 1)
    p.paragraph_format.line_spacing = Pt(1)


# ---------- 日文行解析與聲調 ----------
def split_mora(kana):
    out = []
    for ch in kana:
        if ch in SMALL and out:
            out[-1] += ch
        else:
            out.append(ch)
    return out


def parse_line(line, accent_default):
    items = []
    for raw in line.split("／"):
        raw = raw.strip()
        if not raw:
            continue
        if raw.startswith("!"):
            items.append({"kind": "punct", "text": raw[1:]}); continue
        if raw.startswith("@"):
            items.append({"kind": "label", "text": raw[1:]}); continue
        accent = accent_default
        m = re.search(r"#(null|\d+)$", raw)
        if m:
            accent = None if m.group(1) == "null" else int(m.group(1))
            raw = raw[: m.start()]
        body, *gaps = raw.split("+")
        parts = []  # (base, ruby or None)
        for tok in re.finditer(r"\{([^|}]+)\|([^}]+)\}|([^{]+)", body):
            if tok.group(1):
                parts.append((tok.group(1), tok.group(2)))
            else:
                for mo in split_mora(tok.group(3)):
                    parts.append((mo, None))
        gap_parts = []
        for g in gaps:
            shown, _, said = g.partition("=")
            gap_parts.append((shown, said or shown))
        items.append({"kind": "seg", "parts": parts, "gaps": gap_parts, "accent": accent})
    return items


def pitch_marks(n_body, n_gap, accent):
    """回傳 (body_marks, gap_marks)，每拍為 (high, drop)。accent=None → 全部不畫。"""
    if accent is None:
        return [(False, False)] * n_body, [(False, False)] * n_gap
    body = []
    for i in range(n_body):
        if accent == 0:
            body.append((i >= 1, False))
        elif accent == 1:
            body.append((i == 0, i == 0))
        else:
            body.append((1 <= i <= accent - 1, i == accent - 1))
    gap_high = accent == 0
    return body, [(gap_high, False)] * n_gap


def line_columns(items):
    """展開成欄位：每欄 = dict(ruby, base, high, drop, group)。group 相同的 base 會合併。"""
    cols, gid = [], 0
    for it in items:
        gid += 1
        if it["kind"] in ("punct", "label"):
            cols.append(dict(ruby="", base=it["text"], high=False, drop=False, group=gid,
                             w=SZ_JP * len(it["text"])))
            continue
        morae = []  # (ruby_text, base_text, group)
        for base, ruby in it["parts"]:
            if ruby:
                gid += 1
                rm = split_mora(ruby)
                for k, mo in enumerate(rm):
                    morae.append((mo, base, gid, len(rm), len(base)))
            else:
                gid += 1
                morae.append(("", base, gid, 1, len(base)))
        gap_m = []
        for shown, said in it["gaps"]:
            for mo in split_mora(shown):
                gid += 1
                gap_m.append(("", mo, gid, 1, len(mo)))
        bm, gmk = pitch_marks(len(morae), len(gap_m), it["accent"])
        for (ruby, base, g, span, nbase), (hi, dr) in zip(morae + gap_m, bm + gmk):
            # 欄寬：讀音寬與本文寬取大；合併群組的本文寬平均分攤
            w_ruby = SZ_RUBY * len(ruby) + 2 if ruby else 0
            w_base = (SZ_JP * nbase + 1) / span
            cols.append(dict(ruby=ruby, base=base, high=hi, drop=dr, group=g,
                             w=max(w_ruby, w_base, SZ_JP * 0.9)))
    return cols


def zh_runs(p, zh, arrow=None):
    """藍色中文；arrow 為橘色的語感補充（例：「→口語：朋友間常用」）。"""
    if zh:
        r = p.add_run(zh); set_font(r, FONT_ZH, SZ_ZH, BLUE)
    if arrow:
        r = p.add_run(("　" if zh else "") + "→ " + arrow); set_font(r, FONT_ZH, SZ_ZH, ORANGE)


def add_jp_line(container, line, accent_default, bullet=False, trailing_zh=None, arrow=None, avail_w=None):
    """在 cell 或 document 中加入一行帶讀音與聲調線的日文（嵌套表格）。
    trailing_zh 預設接在日文後面；若給 avail_w 且放不下，就改放到下一行。
    bullet 為 True（•）或字串標籤（例 "2."、"A："）；"補" 會加框。
    在 cell 中會順便處理表格後的空段落（中文放下一行時用它，否則縮到最小）。"""
    items = parse_line(line, accent_default)
    cols = line_columns(items)
    bw = 0
    if bullet:
        btxt = bullet if isinstance(bullet, str) else "•"
        bw = 14 * len(btxt) + 8
        cols.insert(0, dict(ruby="", base=btxt, high=False, drop=False, group=-1, w=bw, boxed=btxt == "補"))
    zh_w = (SZ_ZH * (len(trailing_zh or "") + (len(arrow) + 3 if arrow else 0)) + 18) if (trailing_zh or arrow) else 0
    zh_below = bool(zh_w) and avail_w is not None and sum(c["w"] for c in cols) + zh_w > avail_w
    if zh_w and not zh_below:
        cols.append(dict(ruby="", base=trailing_zh, high=False, drop=False, group=-2,
                         w=zh_w, zh=True, arrow=arrow))
    t = container.add_table(rows=2, cols=len(cols))
    in_cell = hasattr(container, "_tc")
    # 儲存格開頭若只有一個空段落，移除它，讓日文行貼齊上緣
    if in_cell:
        first = container._tc.find(qn("w:p"))
        if first is not None and first.getnext() is t._tbl and not first.xpath(".//w:t"):
            container._tc.remove(first)
    table_borders(t, None)
    fixed_layout(t, [c["w"] for c in cols], {"left": 0, "right": 0, "top": 0, "bottom": 0})
    t.rows[0].height = Pt(11); t.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    t.rows[1].height = Pt(20); t.rows[1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    # 讀音列 + 聲調線
    for j, c in enumerate(cols):
        rc = t.cell(0, j)
        rc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        p = first_par(rc); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if c["ruby"]:
            r = p.add_run(c["ruby"]); set_font(r, FONT_JP, SZ_RUBY)
        sides = {}
        if c["high"]:
            sides["top"] = (PITCH_RED, 6)
        if c["drop"]:
            sides["right"] = (PITCH_RED, 6)
        if sides:
            cell_borders(rc, **sides)
    # 本文列：同一 group 合併
    j = 0
    while j < len(cols):
        k = j
        while k + 1 < len(cols) and cols[k + 1]["group"] == cols[j]["group"]:
            k += 1
        cell = t.cell(1, j) if k == j else t.cell(1, j).merge(t.cell(1, k))
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = first_par(cell)
        c = cols[j]
        if c.get("zh"):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Pt(14)
            zh_runs(p, c["base"], c.get("arrow"))
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(c["base"]); set_font(r, FONT_JP, SZ_JP)
            if c.get("boxed"):
                bdr = OxmlElement("w:bdr")
                for a, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", "000000")):
                    bdr.set(qn(a), v)
                r._element.get_or_add_rPr().append(bdr)
        j = k + 1
    # 表格後的段落
    if zh_below:
        p = container.paragraphs[-1] if in_cell else container.add_paragraph()
        tight(p); p.paragraph_format.left_indent = Pt(bw)
        p.paragraph_format.space_after = Pt(3)
        zh_runs(p, trailing_zh, arrow)
    elif in_cell:
        shrink_trailing(container)
    else:  # 頁面上相鄰的表格會被合併成一張，中間要隔一個極小的段落
        p = container.add_paragraph(); tight(p)
        r = p.add_run(""); set_font(r, FONT_JP, 1)
        p.paragraph_format.line_spacing = Pt(1)
    return t


# ---------- 區塊 ----------
def add_title(doc, text):
    p = doc.add_paragraph(); tight(p, WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(text); set_font(r, FONT_TITLE, 18, bold=True)


def add_tag(doc, text):
    p = doc.add_paragraph(); tight(p, WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(f" {text} "); set_font(r, FONT_TITLE, 16, bold=True)
    rpr = r._element.get_or_add_rPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), TAG_FILL)
    rpr.append(shd)


def outer_table(doc, rows, widths):
    t = doc.add_table(rows=rows, cols=len(widths))
    table_borders(t, BORDER_GRAY, 12)
    fixed_layout(t, widths, {"left": 5, "right": 5, "top": 3, "bottom": 3})
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in t.rows:  # 列不跨頁拆開
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); trPr.append(cs)
    return t


def header_row(t, text, fill=HEAD_FILL):
    cell = t.cell(0, 0).merge(t.cell(0, len(t.columns) - 1))
    shade(cell, fill)
    p = first_par(cell); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); set_font(r, FONT_JP, 12)


def spacer(doc, pt=8):
    p = doc.add_paragraph(); tight(p)
    p.paragraph_format.line_spacing = Pt(pt)


def block_cards(doc, b, acc):
    """詞卡：每列 per_row 格（預設 3；動詞等較長的詞用 2）。
    card 可帶 label（課本編號 "2." 或 "補"），中文放不下時自動換到下一行。"""
    per_row = b.get("per_row", 3)
    w = CONTENT_W / per_row
    cards = b["cards"]
    nrows = (len(cards) + per_row - 1) // per_row
    t = outer_table(doc, nrows, [w] * per_row)
    for idx, card in enumerate(cards):
        cell = t.cell(idx // per_row, idx % per_row)
        first_par(cell)
        add_jp_line(cell, card["jp"], acc, bullet=card.get("label") or False,
                    trailing_zh=card.get("zh"), arrow=card.get("arrow"), avail_w=w - 10)
        add_illust(cell, card.get("illust"))
    spacer(doc)


def block_bullets(doc, b, acc):
    """表格外的例句（範本：詞卡下方的 • 例句）。"""
    for row in b["rows"]:
        add_jp_line(doc, row["jp"], acc, bullet=row.get("bullet", True),
                    trailing_zh=row.get("zh"), arrow=row.get("arrow"), avail_w=CONTENT_W)
    spacer(doc)


def add_hyperlink(par, url, text, size=9):
    part = par.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                         is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(a), FONT_JP)
    col = OxmlElement("w:color"); col.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(size * 2))
    for el in (fonts, col, u, sz):
        rpr.append(el)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    r.append(rpr); r.append(t); h.append(r); par._p.append(h)


def fetch_image(src):
    """src 為本機路徑或網址；網址會下載到 images/ 快取。失敗回傳 None。"""
    if not re.match(r"https?://", src):
        return src if os.path.exists(src) else None
    name = os.path.basename(src.split("?")[0]) or "img"
    path = os.path.join(IMG_CACHE, hashlib.md5(src.encode()).hexdigest()[:8] + "_" + name)
    if os.path.exists(path):
        return path
    try:
        os.makedirs(IMG_CACHE, exist_ok=True)
        req = urllib.request.Request(src, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=20) as resp:
            data = resp.read()
        with open(path, "wb") as f:
            f.write(data)
        return path
    except Exception as e:
        print(f"警告：插圖下載失敗，改用佔位框：{src}（{e}）", file=sys.stderr)
        return None


def add_illust(cell, ill):
    """插圖。ill 可為：None（不放圖）、字串（只有關鍵字）、
    dict(source, title, url[, image])：已選定的圖。有 image（網址或本機路徑）時嵌入圖片，
    下方放來源與可點的連結；沒有 image 或下載失敗時退回文字佔位框。"""
    if not ill:
        return
    p = cell.add_paragraph(); tight(p, WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(24)
    if isinstance(ill, str):
        r = p.add_run(f"[插圖：{ill}]"); set_font(r, FONT_JP, 9, GRAY)
        p.paragraph_format.space_after = Pt(24)
        return
    img = fetch_image(ill["image"]) if ill.get("image") else None
    if img:
        p.paragraph_format.space_before = Pt(6)
        p.add_run().add_picture(img, width=Pt(ILLUST_W))
        p2 = cell.add_paragraph(); tight(p2, WD_ALIGN_PARAGRAPH.CENTER)
        add_hyperlink(p2, ill["url"], ill["source"], 7)
        p2.paragraph_format.space_after = Pt(6)
        return
    r = p.add_run(f"[插圖｜{ill['source']}]"); set_font(r, FONT_JP, 9, GRAY)
    p2 = cell.add_paragraph(); tight(p2, WD_ALIGN_PARAGRAPH.CENTER)
    add_hyperlink(p2, ill["url"], ill["title"], 8)
    p2.paragraph_format.space_after = Pt(24)


def block_sentences(doc, b, acc):
    left = CONTENT_W * 0.72
    rows = b["rows"]
    t = outer_table(doc, len(rows) + 1, [left, CONTENT_W - left])
    header_row(t, b["header"], PRACTICE_FILL if b.get("practice") else HEAD_FILL)
    for i, row in enumerate(rows, start=1):
        lc, rc = t.cell(i, 0), t.cell(i, 1)
        first_par(lc)
        add_jp_line(lc, row["jp"], acc, bullet=row.get("bullet", True))
        rc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        zh_runs(first_par(rc), row["zh"], row.get("arrow"))
    spacer(doc)


def defs_term(lc, row, acc):
    """主題表左欄：詞（或中文詞名）＋下方中文與語感。"""
    shade(lc, HEAD_FILL)
    first_par(lc)
    if "term_jp" in row:
        add_jp_line(lc, row["term_jp"], acc)
    else:
        p = lc.paragraphs[0]; r = p.add_run(row["term_zh"]); set_font(r, FONT_JP, 12)
    if row.get("term_sub"):
        p = lc.add_paragraph(); tight(p); zh_runs(p, row["term_sub"])
    if row.get("arrow"):
        p = lc.add_paragraph(); tight(p); zh_runs(p, None, row["arrow"])


def block_defs(doc, b, acc):
    """主題表。row 有 examples（[{jp, zh, arrow}]）時用三欄：詞｜例句｜中文，
    每個例句一列、左欄跨列合併、同一詞內不畫橫線（範本：擬聲擬態語）；
    否則用兩欄：詞｜中文（zh）或日文清單（jp_list）。"""
    rows = b["rows"]
    if any("examples" in r for r in rows):
        left, right = CONTENT_W * 0.26, CONTENT_W * 0.24
        mid = CONTENT_W - left - right
        n = sum(max(1, len(r.get("examples", []))) for r in rows)
        t = outer_table(doc, n + 1, [left, mid, right])
        header_row(t, b["header"])
        i = 1
        for row in rows:
            exs = row.get("examples") or [{}]
            lc = t.cell(i, 0)
            if len(exs) > 1:
                lc = lc.merge(t.cell(i + len(exs) - 1, 0))
            lc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            defs_term(lc, row, acc)
            for k, ex in enumerate(exs):
                mc, rc = t.cell(i + k, 1), t.cell(i + k, 2)
                cell_borders(mc, right=None); cell_borders(rc, left=None)
                if k > 0:
                    cell_borders(mc, top=None); cell_borders(rc, top=None)
                if k < len(exs) - 1:
                    cell_borders(mc, bottom=None); cell_borders(rc, bottom=None)
                if not ex:
                    continue
                w = sum(c["w"] for c in line_columns(parse_line(ex["jp"], acc))) + 22
                first_par(mc)
                if w > mid - 10:  # 例句太長：例句跨到中文欄，中文放下一行
                    mc = mc.merge(rc)
                    add_jp_line(mc, ex["jp"], acc, bullet=True, trailing_zh=ex.get("zh"),
                                arrow=ex.get("arrow"), avail_w=0)
                else:
                    add_jp_line(mc, ex["jp"], acc, bullet=True)
                    rc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    zh_runs(first_par(rc), ex.get("zh"), ex.get("arrow"))
            i += len(exs)
        spacer(doc)
        return
    left = CONTENT_W * 0.28
    t = outer_table(doc, len(rows) + 1, [left, CONTENT_W - left])
    header_row(t, b["header"])
    for i, row in enumerate(rows, start=1):
        lc, rc = t.cell(i, 0), t.cell(i, 1)
        defs_term(lc, row, acc)
        if "zh" in row:
            rc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            zh_runs(first_par(rc), row["zh"])
        if "jp_list" in row:
            first_par(rc)
            for line in row["jp_list"]:
                add_jp_line(rc, line, acc, bullet=True)
    spacer(doc)


def block_grid(doc, b, acc):
    """對照表：表頭＋欄名列＋資料列。儲存格可為 {"jp": ...}、{"zh": ...} 或字串（當中文）。"""
    ncol = len(b["columns"])
    ratio = b.get("widths") or [1] * ncol
    widths = [CONTENT_W * r / sum(ratio) for r in ratio]
    t = outer_table(doc, len(b["rows"]) + 2, widths)
    header_row(t, b["header"])
    for j, name in enumerate(b["columns"]):
        c = t.cell(1, j); shade(c, HEAD_FILL)
        p = first_par(c); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(name); set_font(r, FONT_JP, 12)
    for i, row in enumerate(b["rows"], start=2):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            first_par(c)
            if isinstance(val, dict) and "jp" in val:
                add_jp_line(c, val["jp"], acc, avail_w=widths[j] - 10)
            else:
                text = val["zh"] if isinstance(val, dict) else val
                p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                zh_runs(p, text)
    spacer(doc)


def block_note(doc, b, acc):
    p = doc.add_paragraph(); tight(p)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("＊　" + b["text"]); set_font(r, FONT_ZH, SZ_ZH, ORANGE)


def block_pending(doc, b, acc):
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph(); tight(p)
    r = p.add_run(b["title"]); set_font(r, FONT_JP, 12, GRAY, bold=True)
    for line in b["lines"]:
        p = doc.add_paragraph(); tight(p); p.paragraph_format.space_after = Pt(3)
        r = p.add_run("・" + line); set_font(r, FONT_JP, 10, GRAY)


BLOCKS = {"cards": block_cards, "bullets": block_bullets, "grid": block_grid,
          "sentences": block_sentences, "defs": block_defs,
          "note": block_note, "pending": block_pending,
          "title": lambda d, b, a: add_title(d, b["text"]),
          "tag": lambda d, b, a: add_tag(d, b["text"])}


def page_setup(doc, head_left, head_right):
    s = doc.sections[0]
    s.page_width, s.page_height = Pt(595.3), Pt(841.9)
    s.left_margin = s.right_margin = Pt(54)
    s.top_margin = Pt(64); s.bottom_margin = Pt(60)
    s.header_distance = Pt(40); s.footer_distance = Pt(50)
    hp = s.header.paragraphs[0]; tight(hp)
    from docx.enum.text import WD_TAB_ALIGNMENT
    ts = hp.paragraph_format.tab_stops
    for pos in (Pt(234), Pt(468)):  # 清掉頁首樣式內建的置中／靠右定位點
        ts.add_tab_stop(pos, WD_TAB_ALIGNMENT.CLEAR)
    ts.add_tab_stop(Pt(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run(f"{head_left}\t{head_right}"); set_font(r, FONT_HEAD, SZ_HEAD, GRAY)
    fp = s.footer.paragraphs[0]; tight(fp, WD_ALIGN_PARAGRAPH.CENTER)
    r = fp.add_run(); set_font(r, FONT_HEAD, SZ_HEAD, GRAY)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        r._element.append(el)


def build(src, dst):
    note = json.load(open(src, encoding="utf-8"))
    doc = Document()
    # 預設段落樣式歸零，避免 Normal 的段後距
    st = doc.styles["Normal"]; st.paragraph_format.space_after = Pt(0)
    page_setup(doc, note["header"]["left"], note["header"]["right"])
    acc = note.get("accent_default", 0)
    for b in note["blocks"]:
        BLOCKS[b["type"]](doc, b, acc)
    doc.save(dst)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
